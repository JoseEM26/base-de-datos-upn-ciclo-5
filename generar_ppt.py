#!/usr/bin/env python3
"""
generar_ppt.py — Generador de la Presentación (PPTX) del Proyecto Final BGG (UPN)
Aplica el sistema de diseño oficial UPN y renderiza diagramas Mermaid vía Chrome headless con CDN.
"""

import os
import re
import subprocess
from PIL import Image, ImageChops

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

# Rutas de entrada y salida
BASE_DIR = "/home/jose/UPN/base-datos-upn"
MD_PATH = os.path.join(BASE_DIR, "docs/proyecto-final/03-entregables/ppt/contenido-ppt.md")
OUT_DIR = os.path.join(BASE_DIR, "docs/proyecto-final/03-entregables/ppt-final")
OUT_PPTX = os.path.join(OUT_DIR, "Presentacion_Proyecto_Final_BGG.pptx")
IMG_DIR = "/tmp/ppt_mermaid_renders"

LOGO_UPN = os.path.join(BASE_DIR, ".claude/skills/upn-document-design/assets/logos/upn-logo.png")
LOGO_BGG = os.path.join(BASE_DIR, ".claude/skills/upn-document-design/assets/logos/bgg-logo.png")

# Colores oficiales UPN
COLOR_GOLD = RGBColor(255, 192, 0)      # #FFC000
COLOR_CREAM = RGBColor(255, 243, 204)   # #FFF3CC
COLOR_BLACK = RGBColor(0, 0, 0)         # #000000
COLOR_WHITE = RGBColor(255, 255, 255)   # #FFFFFF
COLOR_GRAY_DARK = RGBColor(40, 40, 40)
COLOR_GRAY_LIGHT = RGBColor(120, 120, 120)
COLOR_BORDER = RGBColor(220, 220, 220)


def render_mermaid_to_png(mermaid_code: str, output_png: str, width=1600, height=1200):
    """Renderiza diagrama Mermaid a PNG usando Chrome headless y CDN jsdelivr."""
    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<script src="https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.min.js"></script>
</head><body style="margin:0;background:white;">
<div class="mermaid">
{mermaid_code}
</div>
<script>mermaid.initialize({{startOnLoad:true}});</script>
</body></html>"""
    html_path = output_png + ".html"
    with open(html_path, "w", encoding="utf-8") as f:
        f.write(html)
    
    subprocess.run([
        "google-chrome", "--headless", "--disable-gpu", "--no-sandbox",
        f"--screenshot={output_png}",
        f"--window-size={width},{height}",
        "--virtual-time-budget=4000",
        "--default-background-color=FFFFFFFF",
        f"file://{os.path.abspath(html_path)}",
    ], check=True, capture_output=True, timeout=30)

    # Recortar whitespace sobrante
    img = Image.open(output_png).convert("RGB")
    bg = Image.new("RGB", img.size, (255, 255, 255))
    diff = ImageChops.difference(img, bg)
    bbox = diff.getbbox()
    if bbox:
        pad = 20
        l, t, r, b = bbox
        l = max(0, l - pad); t = max(0, t - pad)
        r = min(img.width, r + pad); b = min(img.height, b + pad)
        img.crop((l, t, r, b)).save(output_png)


def add_markdown_runs(paragraph, text, base_font_size=Pt(13), base_color=COLOR_GRAY_DARK):
    """Parsea markdown simple (**negrita** y `codigo`) en runs de texto de python-pptx."""
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        run = paragraph.add_run()
        run.font.name = 'Calibri'
        run.font.size = base_font_size
        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.font.bold = True
            run.font.color.rgb = COLOR_BLACK
        elif token.startswith('`') and token.endswith('`'):
            run.text = token[1:-1]
            run.font.bold = True
            run.font.color.rgb = RGBColor(60, 60, 60)
        else:
            run.text = token
            run.font.bold = False
            run.font.color.rgb = base_color


def apply_master_design(slide, slide_num, total_slides, title_text):
    """Aplica el encabezado y pie de página institucional UPN en diapositivas de contenido."""
    # Título de diapositiva
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(9.8), Inches(0.75))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.name = 'Calibri'
    p.font.size = Pt(21)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLACK

    # Barra dorada bajo el título (#FFC000)
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.22), Inches(11.733), Inches(0.04))
    bar.fill.solid()
    bar.fill.fore_color.rgb = COLOR_GOLD
    bar.line.fill.background()

    # Logo UPN pequeño en esquina superior derecha
    if os.path.exists(LOGO_UPN):
        slide.shapes.add_picture(LOGO_UPN, Inches(11.3), Inches(0.35), height=Inches(0.75))

    # Línea divisoria de pie de página
    f_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(6.85), Inches(11.733), Inches(0.015))
    f_bar.fill.solid()
    f_bar.fill.fore_color.rgb = COLOR_BORDER
    f_bar.line.fill.background()

    # Texto pie de página izquierda
    footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.92), Inches(8.0), Inches(0.35))
    tf_f = footer_box.text_frame
    tf_f.margin_left = tf_f.margin_top = tf_f.margin_right = tf_f.margin_bottom = 0
    p_f = tf_f.paragraphs[0]
    p_f.text = "Curso: Base de Datos — Proyecto Final BGG"
    p_f.font.name = 'Calibri'
    p_f.font.size = Pt(10)
    p_f.font.color.rgb = COLOR_GRAY_LIGHT

    # Numeración de diapositiva a la derecha
    num_box = slide.shapes.add_textbox(Inches(10.5), Inches(6.92), Inches(2.033), Inches(0.35))
    tf_n = num_box.text_frame
    tf_n.margin_left = tf_n.margin_top = tf_n.margin_right = tf_n.margin_bottom = 0
    p_n = tf_n.paragraphs[0]
    p_n.alignment = PP_ALIGN.RIGHT
    p_n.text = f"Diapositiva {slide_num} / {total_slides}"
    p_n.font.name = 'Calibri'
    p_n.font.size = Pt(10)
    p_n.font.color.rgb = COLOR_GRAY_LIGHT


def add_centered_image(slide, image_path, top_pos=Inches(1.4), max_w=Inches(11.733), max_h=Inches(5.2)):
    """Inserta una imagen centrada en la diapositiva respetando su relación de aspecto."""
    if not os.path.exists(image_path):
        print(f"Error: No existe imagen {image_path}")
        return
    im = Image.open(image_path)
    aspect = im.width / im.height

    target_h = max_h
    target_w = target_h * aspect
    if target_w > max_w:
        target_w = max_w
        target_h = target_w / aspect

    slide_w = Inches(13.333)
    left = (slide_w - target_w) / 2
    top = top_pos + (max_h - target_h) / 2
    slide.shapes.add_picture(image_path, left, top, width=target_w, height=target_h)


def create_slide_1_portada(prs):
    """Crea la Diapositiva 1: Portada con diseño oficial UPN y logos UPN + BGG."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])

    # Barra dorada superior
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(0.08))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = COLOR_GOLD
    top_bar.line.fill.background()

    # Logo UPN grande centrado arriba
    if os.path.exists(LOGO_UPN):
        im = Image.open(LOGO_UPN)
        aspect = im.width / im.height
        h = Inches(1.15)
        w = h * aspect
        left = (Inches(13.333) - w) / 2
        slide.shapes.add_picture(LOGO_UPN, left, Inches(0.55), width=w, height=h)

    # Bloque institucional
    inst_box = slide.shapes.add_textbox(Inches(1.5), Inches(1.85), Inches(10.333), Inches(0.6))
    tf = inst_box.text_frame
    tf.word_wrap = True
    p1 = tf.paragraphs[0]
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "UNIVERSIDAD PRIVADA DEL NORTE — FACULTAD DE INGENIERÍA"
    p1.font.name = 'Calibri'
    p1.font.size = Pt(13)
    p1.font.bold = True
    p1.font.color.rgb = COLOR_BLACK

    p2 = tf.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Carrera de Ingeniería de Sistemas Computacionales | Curso: Base de Datos (2026-1)"
    p2.font.name = 'Calibri'
    p2.font.size = Pt(11)
    p2.font.color.rgb = COLOR_GRAY_LIGHT

    # Barra horizontal dorada separadora
    mid_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(2.2), Inches(2.55), Inches(8.933), Inches(0.04))
    mid_bar.fill.solid()
    mid_bar.fill.fore_color.rgb = COLOR_GOLD
    mid_bar.line.fill.background()

    # Título del Proyecto con logo BGG a la izquierda
    if os.path.exists(LOGO_BGG):
        slide.shapes.add_picture(LOGO_BGG, Inches(1.2), Inches(2.75), width=Inches(1.5), height=Inches(1.5))

    title_box = slide.shapes.add_textbox(Inches(2.9), Inches(2.7), Inches(9.2), Inches(1.7))
    tf_t = title_box.text_frame
    tf_t.word_wrap = True
    p_t = tf_t.paragraphs[0]
    p_t.text = "Diseño e Implementación de una Base de Datos Relacional para la Gestión de Microcréditos y Cobranza en Ruta (Caso BGG) aplicando el Patrón PARTY en SQL Server"
    p_t.font.name = 'Calibri'
    p_t.font.size = Pt(18)
    p_t.font.bold = True
    p_t.font.color.rgb = COLOR_BLACK

    # Tarjeta izquierda: Docente y Metodología
    card1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.2), Inches(4.7), Inches(5.2), Inches(2.1))
    card1.fill.solid()
    card1.fill.fore_color.rgb = COLOR_CREAM
    card1.line.color.rgb = COLOR_GOLD
    card1.line.width = Pt(1.5)

    tf_c1 = card1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = tf_c1.margin_top = tf_c1.margin_right = tf_c1.margin_bottom = Inches(0.18)
    
    p = tf_c1.paragraphs[0]
    p.text = "Información del Curso y Proyecto"
    p.font.name = 'Calibri'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLACK

    items_c1 = [
        ("Docente:", "Juan Emilio Asto Vara"),
        ("Sede / Periodo:", "Cajamarca – Perú, 2026-1"),
        ("SGBD:", "Microsoft SQL Server"),
        ("Metodología:", "AS-IS vs. TO-BE (Patrón PARTY)")
    ]
    for lbl, val in items_c1:
        p = tf_c1.add_paragraph()
        r1 = p.add_run()
        r1.text = f"{lbl} "
        r1.font.bold = True
        r1.font.size = Pt(11)
        r1.font.color.rgb = COLOR_BLACK
        r2 = p.add_run()
        r2.text = val
        r2.font.bold = False
        r2.font.size = Pt(11)
        r2.font.color.rgb = COLOR_GRAY_DARK

    # Tarjeta derecha: Integrantes
    card2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.8), Inches(4.7), Inches(5.333), Inches(2.1))
    card2.fill.solid()
    card2.fill.fore_color.rgb = COLOR_WHITE
    card2.line.color.rgb = COLOR_GOLD
    card2.line.width = Pt(1.5)

    tf_c2 = card2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = tf_c2.margin_top = tf_c2.margin_right = tf_c2.margin_bottom = Inches(0.18)

    p = tf_c2.paragraphs[0]
    p.text = "Integrantes (% de Participación)"
    p.font.name = 'Calibri'
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLACK

    integrantes = [
        "Aquino Rivera, Oswaldo Jader (25%)",
        "Espinoza Morales, Jose Angel (25%)",
        "León Ccahuana, Jeffre Carlos (25%)",
        "Ramos Guerra, Jaime Eloy (25%)"
    ]
    for inte in integrantes:
        p = tf_c2.add_paragraph()
        r = p.add_run()
        r.text = f"•  {inte}"
        r.font.size = Pt(11)
        r.font.color.rgb = COLOR_GRAY_DARK


def create_bullets_slide(prs, slide_num, total_slides, title, lines_text):
    """Crea una diapositiva estándar de viñetas con diseño profesional UPN."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_master_design(slide, slide_num, total_slides, title)

    # Contenedor de viñetas
    box = slide.shapes.add_textbox(Inches(0.8), Inches(1.45), Inches(11.733), Inches(5.2))
    tf = box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

    first = True
    for line in lines_text:
        line_clean = line.strip()
        if not line_clean:
            continue
        
        # Determinar nivel de sangría
        is_sub = line.startswith("    - ") or line.startswith("  - ") or line.startswith("\t- ")
        if line_clean.startswith("- "):
            raw_text = line_clean[2:]
        else:
            raw_text = line_clean

        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()

        if is_sub:
            p.level = 1
            p.space_before = Pt(3)
            p.space_after = Pt(3)
            bullet_prefix = "    ▪  "
            add_markdown_runs(p, bullet_prefix + raw_text, base_font_size=Pt(12), base_color=COLOR_GRAY_DARK)
        else:
            p.level = 0
            p.space_before = Pt(6)
            p.space_after = Pt(4)
            bullet_prefix = "•  "
            add_markdown_runs(p, bullet_prefix + raw_text, base_font_size=Pt(13), base_color=COLOR_BLACK)


def create_two_card_slide(prs, slide_num, total_slides, title, card1_title, card1_items, card2_title, card2_items):
    """Crea una diapositiva comparativa con dos tarjetas estilizadas en paralelo."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_master_design(slide, slide_num, total_slides, title)

    card_w = Inches(5.7)
    card_h = Inches(5.1)
    card_y = Inches(1.45)

    # Tarjeta 1
    c1 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), card_y, card_w, card_h)
    c1.fill.solid()
    c1.fill.fore_color.rgb = COLOR_WHITE
    c1.line.color.rgb = COLOR_GOLD
    c1.line.width = Pt(1.5)
    tf1 = c1.text_frame
    tf1.word_wrap = True
    tf1.margin_left = tf1.margin_top = tf1.margin_right = tf1.margin_bottom = Inches(0.25)
    
    p = tf1.paragraphs[0]
    p.text = card1_title
    p.font.name = 'Calibri'
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLACK
    p.space_after = Pt(8)

    for it in card1_items:
        p = tf1.add_paragraph()
        p.space_after = Pt(6)
        is_sub = it.startswith("    - ") or it.startswith("  - ")
        clean_it = it.strip().lstrip("- ").strip()
        prefix = "    ▪  " if is_sub else "•  "
        sz = Pt(11) if is_sub else Pt(12)
        add_markdown_runs(p, prefix + clean_it, base_font_size=sz)

    # Tarjeta 2
    c2 = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.833), card_y, card_w, card_h)
    c2.fill.solid()
    c2.fill.fore_color.rgb = COLOR_CREAM
    c2.line.color.rgb = COLOR_GOLD
    c2.line.width = Pt(1.5)
    tf2 = c2.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_top = tf2.margin_right = tf2.margin_bottom = Inches(0.25)

    p = tf2.paragraphs[0]
    p.text = card2_title
    p.font.name = 'Calibri'
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = COLOR_BLACK
    p.space_after = Pt(8)

    for it in card2_items:
        p = tf2.add_paragraph()
        p.space_after = Pt(6)
        is_sub = it.startswith("    - ") or it.startswith("  - ")
        clean_it = it.strip().lstrip("- ").strip()
        prefix = "    ▪  " if is_sub else "•  "
        sz = Pt(11) if is_sub else Pt(12)
        add_markdown_runs(p, prefix + clean_it, base_font_size=sz)


def create_table_slide_6(prs, slide_num, total_slides):
    """Crea la Diapositiva 6: Matriz de Evaluación de Propuestas con formato cebra UPN."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    apply_master_design(slide, slide_num, total_slides, "Propuestas Evaluadas y Justificación de la Alternativa Elegida")

    table_data = [
        ["Criterio de Evaluación (Escala 0–10)", "Propuesta 1: Modelo Tradicional", "Propuesta 2: Base NoSQL", "Propuesta 3 (Elegida): Patrón PARTY"],
        ["Integridad Transaccional y ACID", "8", "5", "10"],
        ["Eliminación de Redundancia de Datos", "3", "4", "10"],
        ["Trazabilidad Histórica y Auditoría", "4", "6", "10"],
        ["Flexibilidad ante Nuevos Roles", "2", "8", "9"],
        ["Facilidad de Soporte y Madurez SGBD", "8", "6", "9"],
        ["Puntaje Total Ponderado", "25 / 50", "29 / 50", "48 / 50"]
    ]

    rows = len(table_data)
    cols = len(table_data[0])
    table_shape = slide.shapes.add_table(rows, cols, Inches(0.8), Inches(1.45), Inches(11.733), Inches(3.6))
    table = table_shape.table

    table.columns[0].width = Inches(4.333)
    table.columns[1].width = Inches(2.4)
    table.columns[2].width = Inches(2.4)
    table.columns[3].width = Inches(2.6)

    for r_idx, row in enumerate(table_data):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx, c_idx)
            cell.text_frame.word_wrap = True
            cell.text_frame.margin_left = Inches(0.12)
            cell.text_frame.margin_right = Inches(0.12)
            cell.text_frame.margin_top = Inches(0.06)
            cell.text_frame.margin_bottom = Inches(0.06)

            fill = cell.fill
            fill.solid()
            if r_idx == 0:
                fill.fore_color.rgb = COLOR_BLACK
            elif r_idx == rows - 1:
                fill.fore_color.rgb = COLOR_CREAM
            elif r_idx % 2 == 1:
                fill.fore_color.rgb = COLOR_WHITE
            else:
                fill.fore_color.rgb = COLOR_CREAM

            p = cell.text_frame.paragraphs[0]
            p.text = val
            p.font.name = 'Calibri'
            if r_idx == 0:
                p.font.bold = True
                p.font.size = Pt(11)
                p.font.color.rgb = COLOR_WHITE
                if c_idx == 3:
                    p.font.color.rgb = COLOR_GOLD
            elif r_idx == rows - 1:
                p.font.bold = True
                p.font.size = Pt(11)
                p.font.color.rgb = COLOR_BLACK
            else:
                p.font.size = Pt(10)
                p.font.color.rgb = COLOR_BLACK
                if c_idx == 3:
                    p.font.bold = True

            if c_idx > 0:
                p.alignment = PP_ALIGN.CENTER

    # Caja de justificación debajo de la tabla
    callout = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(5.35), Inches(11.733), Inches(1.2))
    callout.fill.solid()
    callout.fill.fore_color.rgb = COLOR_CREAM
    callout.line.color.rgb = COLOR_GOLD
    callout.line.width = Pt(1.5)
    tf_call = callout.text_frame
    tf_call.word_wrap = True
    tf_call.margin_left = tf_call.margin_right = Inches(0.2)
    tf_call.margin_top = Inches(0.12)

    p_c = tf_call.paragraphs[0]
    p_c.text = "Justificación de la Elección Técnica:"
    p_c.font.name = 'Calibri'
    p_c.font.size = Pt(12)
    p_c.font.bold = True
    p_c.font.color.rgb = COLOR_BLACK

    p_c2 = tf_call.add_paragraph()
    p_c2.space_before = Pt(3)
    p_c2.text = "La Propuesta 3 (Patrón PARTY en SQL Server) obtuvo la máxima puntuación gracias a su capacidad de unificar identidades en un núcleo desacoplado, garantizando trazabilidad histórica completa sin duplicar tablas por rol y asegurando consistencia transaccional ACID estricta para préstamos y cobranzas."
    p_c2.font.name = 'Calibri'
    p_c2.font.size = Pt(11)
    p_c2.font.color.rgb = COLOR_GRAY_DARK


def main():
    os.makedirs(OUT_DIR, exist_ok=True)
    os.makedirs(IMG_DIR, exist_ok=True)

    print("1. Leyendo contenido Markdown...")
    with open(MD_PATH, "r", encoding="utf-8") as f:
        md_text = f.read()

    # Extraer bloques de diapositivas
    slide_chunks = re.split(r'\n(?=## Diapositiva \d+:)', md_text)
    slide_chunks = [s for s in slide_chunks if s.strip().startswith("## Diapositiva")]
    total_slides = len(slide_chunks)
    print(f"Total diapositivas identificadas: {total_slides}")

    # Extraer y renderizar diagramas Mermaid
    mermaid_blocks = re.findall(r"```mermaid\n(.*?)```", md_text, re.DOTALL)
    print(f"Total diagramas Mermaid encontrados: {len(mermaid_blocks)}")

    diagram_configs = [
        ("slide4_septe.png", 1600, 1000),
        ("slide7_procesos.png", 1800, 1000),
        ("slide8_party.png", 1600, 1000),
        ("slide9_er.png", 1800, 2400)
    ]

    rendered_images = {}
    for idx, code in enumerate(mermaid_blocks):
        name, w, h = diagram_configs[idx]
        img_path = os.path.join(IMG_DIR, name)
        print(f"Renderizando diagrama {idx+1} ({name})...")
        render_mermaid_to_png(code, img_path, width=w, height=h)
        if os.path.exists(img_path) and os.path.getsize(img_path) > 0:
            sz = os.path.getsize(img_path)
            im = Image.open(img_path)
            print(f"  -> OK: {name} ({im.size[0]}x{im.size[1]} px, {sz} bytes)")
            rendered_images[idx] = img_path
        else:
            raise RuntimeError(f"Error al renderizar diagrama {name}")

    # Inicializar presentación
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    for idx, chunk in enumerate(slide_chunks):
        slide_num = idx + 1
        lines = chunk.strip().split("\n")
        raw_header = lines[0]
        # Extraer título después de "## Diapositiva N: "
        title_match = re.match(r"^## Diapositiva \d+:\s*(.*)$", raw_header)
        title = title_match.group(1).strip() if title_match else raw_header

        print(f"Generando diapositiva {slide_num}: {title}...")

        if slide_num == 1:
            create_slide_1_portada(prs)

        elif slide_num == 2:
            content_lines = [l for l in lines[1:] if l.strip().startswith("- ") or l.strip().startswith("  - ") or l.strip().startswith("    - ")]
            create_bullets_slide(prs, slide_num, total_slides, title, content_lines)

        elif slide_num == 3:
            content_lines = [l for l in lines[1:] if l.strip().startswith("- ") or l.strip().startswith("  - ") or l.strip().startswith("    - ")]
            create_bullets_slide(prs, slide_num, total_slides, title, content_lines)

        elif slide_num == 4:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            apply_master_design(slide, slide_num, total_slides, title)
            add_centered_image(slide, rendered_images[0], top_pos=Inches(1.4), max_w=Inches(11.733), max_h=Inches(5.2))

        elif slide_num == 5:
            content_lines = [l for l in lines[1:] if l.strip().startswith("- ") or l.strip().startswith("  - ") or l.strip().startswith("    - ")]
            create_bullets_slide(prs, slide_num, total_slides, title, content_lines)

        elif slide_num == 6:
            create_table_slide_6(prs, slide_num, total_slides)

        elif slide_num == 7:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            apply_master_design(slide, slide_num, total_slides, title)
            add_centered_image(slide, rendered_images[1], top_pos=Inches(1.4), max_w=Inches(11.733), max_h=Inches(4.4))
            # Callout inferior
            c_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.0), Inches(11.733), Inches(0.6))
            tf = c_box.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            add_markdown_runs(p, "**Impacto Operativo:** Reducción a cero de duplicados de identidad, validación automática de garantías activas y control exacto de cobranza diaria en SQL Server.", base_font_size=Pt(12))

        elif slide_num == 8:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            apply_master_design(slide, slide_num, total_slides, title)
            # Texto superior
            t_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.733), Inches(2.3))
            tf = t_box.text_frame
            tf.word_wrap = True
            tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
            
            p = tf.paragraphs[0]
            add_markdown_runs(p, "•  **Referencia Académica:** Silverston, Len (2001). *The Data Model Resource Book (Vol. 1)*.", base_font_size=Pt(13))
            p.space_after = Pt(4)

            p = tf.add_paragraph()
            add_markdown_runs(p, "•  **Separación de Responsabilidades en Tres Preguntas Ontológicas:**", base_font_size=Pt(13))
            p.space_after = Pt(3)

            q_items = [
                "**1. ¿Quién existe? (Identidad Ontológica):** `PARTY`, `PERSON`, `ORGANIZATION` (inmutables, clave subrogada).",
                "**2. ¿Qué función ejerce? (Rol de Negocio):** `PARTY_ROLE` (`CLIENTE`, `PRESTATARIO`, `AVAL`, `COBRADOR`) con vigencia temporal.",
                "**3. ¿Con quién se relaciona? (Vínculo Dirigido):** `PARTY_RELATIONSHIP` (vínculos dirigidos de aval y asignación de cobranza)."
            ]
            for q in q_items:
                p = tf.add_paragraph()
                p.space_after = Pt(2)
                add_markdown_runs(p, f"    ▪  {q}", base_font_size=Pt(11.5))

            # Diagrama inferior
            add_centered_image(slide, rendered_images[2], top_pos=Inches(3.85), max_w=Inches(11.733), max_h=Inches(2.8))

        elif slide_num == 9:
            slide = prs.slides.add_slide(prs.slide_layouts[6])
            apply_master_design(slide, slide_num, total_slides, title)
            add_centered_image(slide, rendered_images[3], top_pos=Inches(1.35), max_w=Inches(11.733), max_h=Inches(5.35))

        elif slide_num == 10:
            c1_items = [
                "Tablas aisladas y redundantes por rol (`tbl_clientes`, `tbl_avales`, `tbl_cobradores`).",
                "Columnas NULL masivas (razón social nula para personas, nombres/género nulos para empresas).",
                "Atributos multivalor repetitivos para números telefónicos y mecanismos de contacto.",
                "Uso de documentos (DNI/RUC) como PK: falla con extranjeros o clientes sin documento definitivo.",
                "Anomalías de actualización: modificar un teléfono deja las demás tablas desactualizadas."
            ]
            c2_items = [
                "**1FN (Atomicidad e Identidad Subrogada):** Claves `BIGINT` únicas (`party_id`), atributos atómicos y teléfonos desacoplados en `CONTACT_MECHANISM`.",
                "**2FN (Dependencia Funcional Completa):** Separación de herencia 1:1 (`PERSON` y `ORGANIZATION`) dependiendo exclusivamente de su clave.",
                "**3FN / BCNF (Cero Dependencias Transitivas):** Eliminación total de anomalías; los roles y créditos cuelgan de claves especializadas sin contaminar la identidad raíz."
            ]
            create_two_card_slide(prs, slide_num, total_slides, title, "Estado AS-IS (No Normalizado / 1FN Deficiente)", c1_items, "Estado TO-BE (Normalizado en 3FN y BCNF)", c2_items)

        elif slide_num == 11:
            content_lines = [l for l in lines[1:] if l.strip().startswith("- ") or l.strip().startswith("  - ") or l.strip().startswith("    - ")]
            create_bullets_slide(prs, slide_num, total_slides, title, content_lines)

        elif slide_num == 12:
            content_lines = [l for l in lines[1:] if l.strip().startswith("- ") or l.strip().startswith("  - ") or l.strip().startswith("    - ")]
            create_bullets_slide(prs, slide_num, total_slides, title, content_lines)

        elif slide_num == 13:
            c1_items = [
                "**Backup Completo (Full):** Ejecución semanal programada (domingos a las 00:00 hrs) capturando la totalidad de la base de datos.",
                "**Backup Diferencial:** Ejecución diaria de lunes a sábado (23:00 hrs) respaldando únicamente páginas modificadas desde el último Full.",
                "**Backup de Log de Transacciones:** Ejecución periódica cada 1 hora en horario operativo (07:00 a 20:00 hrs) bajo Recovery Model `FULL`."
            ]
            c2_items = [
                "**RPO (Punto Objetivo de Recuperación):** Menor a 1 hora de pérdida potencial de transacciones ante caída catastrófica.",
                "**RTO (Tiempo Objetivo de Recuperación):** Restauración completa del servicio en menos de 30 minutos.",
                "**Validación Periódica:** Scripts automáticos de verificación de integridad de respaldo con `RESTORE VERIFYONLY`.",
                "**Almacenamiento Fuera de Sitio:** Respaldo secundario cifrado para continuidad del negocio."
            ]
            create_two_card_slide(prs, slide_num, total_slides, title, "Estrategia de Respaldos (Regla 3-2-1 en SQL Server)", c1_items, "Métricas de Continuidad y Recuperación (RPO / RTO)", c2_items)

        elif slide_num == 14:
            content_lines = [l for l in lines[1:] if l.strip().startswith("- ") or l.strip().startswith("  - ") or l.strip().startswith("    - ")]
            create_bullets_slide(prs, slide_num, total_slides, title, content_lines)

        elif slide_num == 15:
            c1_items = [
                "El patrón PARTY eliminó de raíz la redundancia de identidades y la pérdida de trazabilidad que aquejaban al modelo AS-IS de BGG.",
                "La arquitectura en Microsoft SQL Server garantiza consistencia transaccional ACID estricta para préstamos y liquidaciones de cobranza en ruta.",
                "El modelado de vínculos (`PARTY_RELATIONSHIP`) permite rastrear avales activos y reasignaciones de cobradores sin alterar el esquema físico relacional."
            ]
            c2_items = [
                "**Vistas Indexadas (Indexed Views):** Implementar vistas materializadas en SQL Server para acelerar consultas analíticas frecuentes de cobranza diaria.",
                "**Triggers de Validación Temporal:** Incorporar restricciones procedurales para impedir el solapamiento de vigencias en `PARTY_ROLE`.",
                "**Automatización de Respaldos:** Configurar alertas automatizadas en SQL Server Agent para notificar desviaciones en el plan de backups."
            ]
            create_two_card_slide(prs, slide_num, total_slides, title, "Conclusiones Técnicas del Proyecto", c1_items, "Recomendaciones de Ingeniería y Escalabilidad", c2_items)

    print(f"Guardando presentación en {OUT_PPTX}...")
    prs.save(OUT_PPTX)
    print("¡Presentación guardada exitosamente!")


if __name__ == "__main__":
    main()
