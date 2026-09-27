#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
generar_informe.py
Genera el informe oficial en formato Word (.docx) aplicando el sistema de diseño visual de la UPN.
Convierte los 8 archivos Markdown en un documento integrado con estilos institucionales,
tablas con formato cebra, portada institucional, header/footer con numeración de página
y diagramas Mermaid renderizados a imágenes PNG vía Chrome headless.
"""

import os
import re
import glob
import subprocess
from PIL import Image, ImageChops

import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

BASE_DIR = "/home/jose/UPN/base-datos-upn"
INFORME_DIR = os.path.join(BASE_DIR, "docs/proyecto-final/03-entregables/informe")
OUTPUT_DIR = os.path.join(BASE_DIR, "docs/proyecto-final/03-entregables/informe-final")
TMP_IMG_DIR = os.path.join(BASE_DIR, "scripts/tmp_mermaid")
LOGO_UPN = os.path.join(BASE_DIR, ".claude/skills/upn-document-design/assets/logos/upn-logo.png")
LOGO_BGG = os.path.join(BASE_DIR, ".claude/skills/upn-document-design/assets/logos/bgg-logo.png")

DOCX_OUTPUT = os.path.join(OUTPUT_DIR, "Informe_Proyecto_Final_BGG.docx")

# Paleta UPN
COLOR_GOLD = "FFC000"
COLOR_CREAM = "FFF3CC"
COLOR_BLACK = "000000"
COLOR_WHITE = "FFFFFF"
COLOR_GRAY_LIGHT = "F4F4F4"
COLOR_GRAY_BORDER = "CCCCCC"
COLOR_TEXT_DARK = "111111"

RGB_GOLD = RGBColor(0xFF, 0xC0, 0x00)
RGB_BLACK = RGBColor(0x00, 0x00, 0x00)
RGB_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
RGB_DARK = RGBColor(0x22, 0x22, 0x22)
RGB_MUTED = RGBColor(0x66, 0x66, 0x66)

os.makedirs(OUTPUT_DIR, exist_ok=True)
os.makedirs(TMP_IMG_DIR, exist_ok=True)

# ---------------------------------------------------------------------------
# XML Helpers para python-docx
# ---------------------------------------------------------------------------

def set_cell_background(cell, fill_hex):
    """Asigna color de fondo a una celda de tabla."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('shd'):
            tcPr.remove(child)
    ns = nsdecls('w')
    shd = parse_xml(f'<w:shd {ns} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    """Establece padding interno de celda en dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcMar'):
            tcPr.remove(child)
    ns = nsdecls('w')
    tcMar = parse_xml(f'<w:tcMar {ns}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_table_borders(table, color="CCCCCC", sz="4", val="single"):
    """Establece bordes limpios para una tabla completa."""
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    ns = nsdecls('w')
    tblBorders = parse_xml(f'''
    <w:tblBorders {ns}>
        <w:top w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:left w:val="none"/>
        <w:bottom w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:right w:val="none"/>
        <w:insideH w:val="{val}" w:sz="{sz}" w:space="0" w:color="{color}"/>
        <w:insideV w:val="none"/>
    </w:tblBorders>
    ''')
    tblPr.append(tblBorders)

def add_golden_divider(paragraph):
    """Agrega un borde inferior dorado (#FFC000) a un párrafo."""
    pPr = paragraph._p.get_or_add_pPr()
    for child in list(pPr):
        if child.tag.endswith('pBdr'):
            pPr.remove(child)
    ns = nsdecls('w')
    pBdr = parse_xml(f'''
    <w:pBdr {ns}>
        <w:bottom w:val="single" w:sz="18" w:space="4" w:color="{COLOR_GOLD}"/>
    </w:pBdr>
    ''')
    pPr.append(pBdr)

def add_page_number_field(run):
    """Inserta el campo dinámico PAGE en un run de texto."""
    fldSimple = OxmlElement('w:fldSimple')
    fldSimple.set(qn('w:instr'), 'PAGE')
    run._r.append(fldSimple)

# ---------------------------------------------------------------------------
# Renderizado de Mermaid a PNG vía Chrome Headless
# ---------------------------------------------------------------------------

mermaid_counter = 0

def render_mermaid_to_png(mermaid_code: str, output_png: str) -> bool:
    """Renderiza código Mermaid a PNG utilizando Chrome headless y CDN jsDelivr."""
    try:
        lines = mermaid_code.strip().splitlines()
        num_lines = len(lines)
        width = 1600
        height = max(1400, num_lines * 45)
        
        # Ajuste de dimensiones según tipo
        if "gantt" in mermaid_code.lower():
            width = 1800
            height = 1000
        elif "flowchart" in mermaid_code.lower() or "graph" in mermaid_code.lower():
            if num_lines > 12:
                height = 2400
                width = 1600
        
        html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<script src="https://cdn.jsdelivr.net/npm/mermaid@10/dist/mermaid.min.js"></script>
</head><body style="margin:0;background:white;">
<div class="mermaid">
{mermaid_code}
</div>
<script>mermaid.initialize({{startOnLoad:true, theme: "neutral", securityLevel: "loose"}});</script>
</body></html>"""
        
        html_path = output_png + ".html"
        with open(html_path, "w", encoding="utf-8") as f:
            f.write(html)
            
        cmd = [
            "google-chrome", "--headless", "--disable-gpu", "--no-sandbox",
            f"--screenshot={output_png}",
            f"--window-size={width},{height}",
            "--virtual-time-budget=5000",
            "--default-background-color=FFFFFFFF",
            f"file://{html_path}",
        ]
        subprocess.run(cmd, check=True, capture_output=True, timeout=35)
        
        if os.path.exists(output_png) and os.path.getsize(output_png) > 100:
            # Recortar whitespace sobrante
            img = Image.open(output_png).convert("RGB")
            bg = Image.new("RGB", img.size, (255, 255, 255))
            diff = ImageChops.difference(img, bg)
            bbox = diff.getbbox()
            if bbox:
                pad = 25
                l, t, r, b = bbox
                l = max(0, l - pad)
                t = max(0, t - pad)
                r = min(img.width, r + pad)
                b = min(img.height, b + pad)
                cropped = img.crop((l, t, r, b))
                cropped.save(output_png)
            print(f"[OK] Diagrama Mermaid renderizado en {output_png} ({os.path.getsize(output_png)} bytes)")
            return True
        else:
            print(f"[ERROR] Archivo generado vacío o nulo para {output_png}")
            return False
    except Exception as e:
        print(f"[ERROR] Excepción renderizando Mermaid a PNG: {e}")
        return False

# ---------------------------------------------------------------------------
# Creación y Configuración del Documento Word
# ---------------------------------------------------------------------------

def crear_documento():
    doc = docx.Document()
    
    # Configuración de página: A4, márgenes estándar 2.5 cm (aprox 1 pulgada)
    for section in doc.sections:
        section.page_width = Inches(8.27)   # A4
        section.page_height = Inches(11.69)
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
        section.different_first_page_header_footer = True
    
    # Configurar estilo Normal
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Calibri'
    style_normal.font.size = Pt(11)
    style_normal.font.color.rgb = RGB_BLACK
    style_normal.paragraph_format.line_spacing = 1.15
    style_normal.paragraph_format.space_after = Pt(6)
    
    return doc

def configurar_header_footer(doc):
    """Configura encabezado y pie de página institucionales para las páginas de contenido."""
    section = doc.sections[0]
    
    # Encabezado (Header)
    header = section.header
    p_head = header.paragraphs[0]
    p_head.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_head.paragraph_format.space_after = Pt(4)
    if os.path.exists(LOGO_UPN):
        run_logo = p_head.add_run()
        run_logo.add_picture(LOGO_UPN, width=Inches(1.3))
    add_golden_divider(p_head)
    
    # Pie de página (Footer)
    footer = section.footer
    tbl_foot = footer.add_table(rows=1, cols=2, width=Inches(6.27))
    tbl_foot.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_foot, color=COLOR_GOLD, sz="12", val="single") # Barra dorada superior
    
    c_left = tbl_foot.cell(0, 0)
    c_right = tbl_foot.cell(0, 1)
    c_left.width = Inches(4.5)
    c_right.width = Inches(1.77)
    
    set_cell_margins(c_left, top=80, bottom=40, left=0, right=0)
    set_cell_margins(c_right, top=80, bottom=40, left=0, right=0)
    
    p_fl = c_left.paragraphs[0]
    p_fl.paragraph_format.space_after = Pt(0)
    r_fl = p_fl.add_run("Curso: Base de Datos — Informe Proyecto Final (Caso BGG)")
    r_fl.font.name = "Calibri"
    r_fl.font.size = Pt(8.5)
    r_fl.font.color.rgb = RGB_MUTED
    
    p_fr = c_right.paragraphs[0]
    p_fr.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_fr.paragraph_format.space_after = Pt(0)
    r_fr = p_fr.add_run("Página ")
    r_fr.font.name = "Calibri"
    r_fr.font.size = Pt(8.5)
    r_fr.font.color.rgb = RGB_MUTED
    add_page_number_field(r_fr)

def agregar_portada(doc):
    """Genera la portada institucional UPN conforme a las directrices oficiales."""
    p_top = doc.add_paragraph()
    p_top.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_top.paragraph_format.space_before = Pt(10)
    p_top.paragraph_format.space_after = Pt(20)
    
    if os.path.exists(LOGO_UPN):
        run_logo = p_top.add_run()
        run_logo.add_picture(LOGO_UPN, width=Inches(2.5))
    
    # Caja Institucional Negra con acento dorado
    tbl_box = doc.add_table(rows=1, cols=1)
    tbl_box.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_box.autofit = False
    
    cell_box = tbl_box.cell(0, 0)
    cell_box.width = Inches(6.27)
    set_cell_background(cell_box, COLOR_BLACK)
    set_cell_margins(cell_box, top=200, bottom=200, left=240, right=240)
    
    p_inst = cell_box.paragraphs[0]
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.line_spacing = 1.2
    p_inst.paragraph_format.space_after = Pt(0)
    
    r1 = p_inst.add_run("UNIVERSIDAD PRIVADA DEL NORTE\n")
    r1.font.name = "Calibri"
    r1.font.size = Pt(13)
    r1.font.bold = True
    r1.font.color.rgb = RGB_WHITE
    
    r2 = p_inst.add_run("FACULTAD DE INGENIERÍA\n")
    r2.font.name = "Calibri"
    r2.font.size = Pt(12)
    r2.font.bold = True
    r2.font.color.rgb = RGB_GOLD
    
    r3 = p_inst.add_run("CARRERA DE INGENIERÍA DE SISTEMAS COMPUTACIONALES\n")
    r3.font.name = "Calibri"
    r3.font.size = Pt(11)
    r3.font.bold = True
    r3.font.color.rgb = RGB_WHITE
    
    r4 = p_inst.add_run("CURSO: BASE DE DATOS")
    r4.font.name = "Calibri"
    r4.font.size = Pt(11)
    r4.font.bold = True
    r4.font.color.rgb = RGB_GOLD
    
    # Espaciado
    p_sp1 = doc.add_paragraph()
    p_sp1.paragraph_format.space_before = Pt(30)
    p_sp1.paragraph_format.space_after = Pt(10)
    
    # Título del Proyecto
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_after = Pt(10)
    
    r_tit_label = p_title.add_run("INFORME DEL PROYECTO FINAL\n")
    r_tit_label.font.name = "Calibri"
    r_tit_label.font.size = Pt(13)
    r_tit_label.font.bold = True
    r_tit_label.font.color.rgb = RGB_MUTED
    
    r_title = p_title.add_run(
        "DISEÑO E IMPLEMENTACIÓN DE UNA BASE DE DATOS RELACIONAL BAJO EL PATRÓN DE MODELADO PARTY "
        "PARA LA GESTIÓN DE IDENTIDAD CENTRALIZADA, ROLES DINÁMICOS Y TRAZABILIDAD DE COBRANZA "
        "EN LA ENTIDAD DE MICROCRÉDITOS BGG"
    )
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(16)
    r_title.font.bold = True
    r_title.font.color.rgb = RGB_BLACK
    add_golden_divider(p_title)
    
    # Espaciado
    p_sp2 = doc.add_paragraph()
    p_sp2.paragraph_format.space_before = Pt(25)
    p_sp2.paragraph_format.space_after = Pt(5)
    
    # Tabla de Integrantes y Participación
    p_aut_tit = doc.add_paragraph()
    p_aut_tit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_aut_tit.paragraph_format.space_after = Pt(6)
    r_at = p_aut_tit.add_run("AUTORES (ORDEN ALFABÉTICO):")
    r_at.font.name = "Calibri"
    r_at.font.size = Pt(10.5)
    r_at.font.bold = True
    r_at.font.color.rgb = RGB_BLACK
    
    tbl_autores = doc.add_table(rows=5, cols=3)
    tbl_autores.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl_autores, color=COLOR_GOLD, sz="6", val="single")
    
    headers = ["Apellidos y Nombres", "Código UPN", "% Participación"]
    for j, h in enumerate(headers):
        cell = tbl_autores.cell(0, j)
        set_cell_background(cell, COLOR_BLACK)
        set_cell_margins(cell, top=80, bottom=80, left=120, right=120)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(h)
        r.font.name = "Calibri"
        r.font.size = Pt(9.5)
        r.font.bold = True
        r.font.color.rgb = RGB_WHITE
    
    integrantes = [
        ("Aquino Rivera, Oswaldo Jader", "N00571142", "25%"),
        ("Espinoza Morales, Jose Angel", "N00575318", "25%"),
        ("León Ccahuana, Jeffre Carlos", "N00423806", "25%"),
        ("Ramos Guerra, Jaime Eloy", "N00377672", "25%"),
    ]
    
    for i, row in enumerate(integrantes, start=1):
        bg = COLOR_WHITE if i % 2 != 0 else COLOR_CREAM
        for j, val in enumerate(row):
            cell = tbl_autores.cell(i, j)
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=120, right=120)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if j > 0 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.name = "Calibri"
            r.font.size = Pt(9.5)
            r.font.bold = (j == 2)
    
    # Docente y Sede
    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_before = Pt(35)
    p_meta.paragraph_format.space_after = Pt(0)
    
    r_doc = p_meta.add_run("Docente: Ing. Juan Emilio Asto Vara\n")
    r_doc.font.name = "Calibri"
    r_doc.font.size = Pt(11)
    r_doc.font.bold = True
    
    r_sede = p_meta.add_run("Cajamarca – Perú, 2026-1")
    r_sede.font.name = "Calibri"
    r_sede.font.size = Pt(10.5)
    r_sede.font.color.rgb = RGB_MUTED
    
    # Salto de página para separar la portada del contenido
    doc.add_page_break()

# ---------------------------------------------------------------------------
# Parser y Convertidor de Elementos Markdown
# ---------------------------------------------------------------------------

def formatear_runs_en_parrafo(p, texto, default_font="Calibri", default_size=11, default_color=None):
    """Convierte texto con markdown inline (**negrita**, *cursiva*, `código`) en runs con formato."""
    p.text = "" # limpiar
    # Tokenizer simple por regex
    tokens = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`|\$.*?\$)', texto)
    for token in tokens:
        if not token:
            continue
        run = p.add_run()
        run.font.name = default_font
        run.font.size = Pt(default_size)
        if default_color:
            run.font.color.rgb = default_color
            
        if token.startswith('**') and token.endswith('**'):
            run.text = token[2:-2]
            run.font.bold = True
        elif token.startswith('*') and token.endswith('*') and not token.startswith('**'):
            run.text = token[1:-1]
            run.font.italic = True
        elif token.startswith('`') and token.endswith('`'):
            run.text = token[1:-1]
            run.font.name = "Consolas"
            run.font.size = Pt(default_size - 1)
        elif token.startswith('$') and token.endswith('$'):
            run.text = token[1:-1]
            run.font.italic = True
        else:
            run.text = token

def agregar_encabezado_seccion(doc, texto, nivel):
    p = doc.add_paragraph()
    p.paragraph_format.keep_with_next = True
    texto_limpio = texto.strip().lstrip('#').strip()
    
    if nivel == 1:
        p.paragraph_format.space_before = Pt(22)
        p.paragraph_format.space_after = Pt(8)
        run = p.add_run(texto_limpio.upper())
        run.font.name = 'Calibri'
        run.font.size = Pt(15)
        run.font.bold = True
        run.font.color.rgb = RGB_BLACK
        add_golden_divider(p)
    elif nivel == 2:
        p.paragraph_format.space_before = Pt(16)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run(texto_limpio.upper())
        run.font.name = 'Calibri'
        run.font.size = Pt(13)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    elif nivel == 3:
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(texto_limpio)
        run.font.name = 'Calibri'
        run.font.size = Pt(11.5)
        run.font.bold = True
        run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    else:
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(texto_limpio)
        run.font.name = 'Calibri'
        run.font.size = Pt(11)
        run.font.bold = True
        run.font.italic = True

def agregar_tabla_markdown(doc, lineas_tabla):
    """Convierte líneas de tabla Markdown en una tabla Word estilizada con formato cebra institucional."""
    filas = []
    for l in lineas_tabla:
        l = l.strip()
        if not l.startswith('|'):
            continue
        partes = [c.strip() for c in l.strip('|').split('|')]
        # Ignorar línea divisoria |---|---|
        if any(set(c).issubset({'-', ':', ' '}) for c in partes if c):
            continue
        filas.append(partes)
    
    if not filas or len(filas) < 1:
        return
    
    num_cols = max(len(r) for r in filas)
    tbl = doc.add_table(rows=len(filas), cols=num_cols)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(tbl, color=COLOR_GOLD, sz="6", val="single")
    
    for i, fila in enumerate(filas):
        es_header = (i == 0)
        bg = COLOR_BLACK if es_header else (COLOR_WHITE if i % 2 != 0 else COLOR_CREAM)
        
        for j in range(num_cols):
            val = fila[j] if j < len(fila) else ""
            cell = tbl.cell(i, j)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_background(cell, bg)
            set_cell_margins(cell, top=70, bottom=70, left=100, right=100)
            
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            
            if es_header:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(val)
                r.font.name = 'Calibri'
                r.font.size = Pt(9.5)
                r.font.bold = True
                r.font.color.rgb = RGB_WHITE
            else:
                formatear_runs_en_parrafo(p, val, default_font="Calibri", default_size=9.5)
    
    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_before = Pt(4)
    p_post.paragraph_format.space_after = Pt(8)

def agregar_bloque_codigo(doc, codigo, lenguaje=""):
    """Inserta bloque de código estilizado en caja con fondo gris y fuente Consolas."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_background(cell, COLOR_GRAY_LIGHT)
    set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
    
    # Borde sutil
    tblPr = tbl._tbl.tblPr
    ns = nsdecls('w')
    tblBorders = parse_xml(f'''
    <w:tblBorders {ns}>
        <w:top w:val="single" w:sz="4" w:space="0" w:color="{COLOR_GRAY_BORDER}"/>
        <w:left w:val="single" w:sz="12" w:space="0" w:color="{COLOR_GOLD}"/>
        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{COLOR_GRAY_BORDER}"/>
        <w:right w:val="single" w:sz="4" w:space="0" w:color="{COLOR_GRAY_BORDER}"/>
    </w:tblBorders>
    ''')
    tblPr.append(tblBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.0
    
    # Dividir líneas de código
    for idx_l, line in enumerate(codigo.splitlines()):
        if idx_l > 0:
            p = cell.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.0
        run = p.add_run(line if line else " ")
        run.font.name = "Consolas"
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
        
    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(6)

def agregar_callout(doc, texto):
    """Inserta una caja de nota/resumen con fondo crema claro y barra izquierda dorada."""
    tbl = doc.add_table(rows=1, cols=1)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    
    cell = tbl.cell(0, 0)
    cell.width = Inches(6.27)
    set_cell_background(cell, COLOR_CREAM)
    set_cell_margins(cell, top=100, bottom=100, left=160, right=140)
    
    tblPr = tbl._tbl.tblPr
    ns = nsdecls('w')
    tblBorders = parse_xml(f'''
    <w:tblBorders {ns}>
        <w:top w:val="none"/>
        <w:left w:val="single" w:sz="24" w:space="0" w:color="{COLOR_GOLD}"/>
        <w:bottom w:val="none"/>
        <w:right w:val="none"/>
    </w:tblBorders>
    ''')
    tblPr.append(tblBorders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    formatear_runs_en_parrafo(p, texto, default_font="Calibri", default_size=10, default_color=RGB_DARK)
    
    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(6)

def procesar_archivo_md(doc, ruta_md, es_primer_archivo=False):
    """Lee y procesa un archivo Markdown, convirtiendo cada bloque a elementos docx."""
    global mermaid_counter
    print(f"Procesando archivo: {os.path.basename(ruta_md)}...")
    
    with open(ruta_md, "r", encoding="utf-8") as f:
        contenido = f.read()
    
    # Si es el primer archivo (00-caratula-indices-resumen.md), omitir la carátula textual
    # porque ya tenemos la portada oficial UPN generada al inicio
    if es_primer_archivo:
        # Encontrar dónde empieza "## Estructura de Índices" o "## Resumen"
        match = re.search(r'(##\s*Estructura de Índices.*)', contenido, re.DOTALL)
        if match:
            contenido = match.group(1)
        else:
            # Si no, quitar la sección carátula inicial
            contenido = re.sub(r'#.*?(?=##\s*Resumen|##\s*Estructura)', '', contenido, flags=re.DOTALL)
    
    lineas = contenido.splitlines()
    i = 0
    n = len(lineas)
    
    while i < n:
        linea = lineas[i]
        linea_strip = linea.strip()
        
        # 1. Bloque de Código (```)
        if linea_strip.startswith("```"):
            lenguaje = linea_strip.lstrip("`").strip().lower()
            bloque_codigo = []
            i += 1
            while i < n and not lineas[i].strip().startswith("```"):
                bloque_codigo.append(lineas[i])
                i += 1
            i += 1 # saltar cierre ```
            codigo_str = "\n".join(bloque_codigo)
            
            if lenguaje == "mermaid":
                mermaid_counter += 1
                img_path = os.path.join(TMP_IMG_DIR, f"diagram_{mermaid_counter}.png")
                exito = render_mermaid_to_png(codigo_str, img_path)
                
                if exito and os.path.exists(img_path):
                    p_img = doc.add_paragraph()
                    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_img.paragraph_format.space_before = Pt(8)
                    p_img.paragraph_format.space_after = Pt(6)
                    
                    # Calcular ancho óptimo
                    img_pil = Image.open(img_path)
                    w_px, h_px = img_pil.size
                    ratio = h_px / w_px
                    max_w = 6.0
                    ancho_in = min(max_w, 6.0)
                    if ratio > 1.3:
                        ancho_in = min(4.8, 6.0)
                    
                    run_img = p_img.add_run()
                    run_img.add_picture(img_path, width=Inches(ancho_in))
                    
                    p_cap = doc.add_paragraph()
                    p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p_cap.paragraph_format.space_after = Pt(10)
                    r_cap = p_cap.add_run(f"Figura {mermaid_counter}: Diagrama del sistema (Renderizado oficial)")
                    r_cap.font.name = "Calibri"
                    r_cap.font.size = Pt(9)
                    r_cap.font.italic = True
                    r_cap.font.color.rgb = RGB_MUTED
                else:
                    # Fallback si falla renderizado: mostrar como código
                    agregar_bloque_codigo(doc, codigo_str, lenguaje="mermaid")
            else:
                agregar_bloque_codigo(doc, codigo_str, lenguaje=lenguaje)
            continue
            
        # 2. Tablas Markdown (|...|)
        if linea_strip.startswith("|") and linea_strip.endswith("|"):
            lineas_tabla = []
            while i < n and lineas[i].strip().startswith("|") and lineas[i].strip().endswith("|"):
                lineas_tabla.append(lineas[i])
                i += 1
            agregar_tabla_markdown(doc, lineas_tabla)
            continue
            
        # 3. Encabezados (#, ##, ###, ####)
        if linea_strip.startswith("#"):
            nivel = 0
            while nivel < len(linea_strip) and linea_strip[nivel] == '#':
                nivel += 1
            agregar_encabezado_seccion(doc, linea_strip, nivel)
            i += 1
            continue
            
        # 4. Citas y Callouts (> ...)
        if linea_strip.startswith(">"):
            bloque_callout = []
            while i < n and lineas[i].strip().startswith(">"):
                bloque_callout.append(lineas[i].strip().lstrip(">").strip())
                i += 1
            texto_callout = " ".join(bloque_callout)
            # Limpiar marcadores de tipo [!NOTE] o [!IMPORTANT]
            texto_callout = re.sub(r'\[!(NOTE|IMPORTANT|TIP|WARNING|CAUTION)\]', '', texto_callout).strip()
            agregar_callout(doc, texto_callout)
            continue
            
        # 5. Listas con Viñetas (- o * )
        if re.match(r'^[\*\-]\s+', linea_strip):
            texto_item = re.sub(r'^[\*\-]\s+', '', linea_strip)
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            formatear_runs_en_parrafo(p, texto_item)
            i += 1
            continue
            
        # 6. Listas Numeradas (1. 2. etc.)
        if re.match(r'^\d+\.\s+', linea_strip):
            texto_item = re.sub(r'^\d+\.\s+', '', linea_strip)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            formatear_runs_en_parrafo(p, texto_item)
            i += 1
            continue
            
        # 7. Separadores horizontales (--- o ***)
        if re.match(r'^(-{3,}|\*{3,})$', linea_strip):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            add_golden_divider(p)
            i += 1
            continue
            
        # 8. Párrafo Normal o línea en blanco
        if linea_strip:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.line_spacing = 1.15
            formatear_runs_en_parrafo(p, linea_strip)
            
        i += 1

# ---------------------------------------------------------------------------
# Función Principal
# ---------------------------------------------------------------------------

def main():
    print("=" * 70)
    print("INICIANDO GENERACIÓN DEL INFORME OFICIAL UPN (.DOCX)")
    print("=" * 70)
    
    doc = crear_documento()
    
    # 1. Portada Institucional
    print("Generando portada institucional UPN...")
    agregar_portada(doc)
    
    # 2. Configurar Encabezado y Pie de Página
    print("Configurando encabezados y pies de página...")
    configurar_header_footer(doc)
    
    # 3. Leer los 8 archivos Markdown en orden secuencial
    archivos_md = [
        "00-caratula-indices-resumen.md",
        "01-introduccion.md",
        "02-marco-teorico.md",
        "03-herramientas-ingenieria.md",
        "04-generacion-soluciones.md",
        "05-metodologia-desarrollo.md",
        "06-conclusiones-recomendaciones.md",
        "07-referencias-anexos.md",
    ]
    
    for idx, nombre_archivo in enumerate(archivos_md):
        ruta_completa = os.path.join(INFORME_DIR, nombre_archivo)
        if not os.path.exists(ruta_completa):
            print(f"[ERROR CRÍTICO] Archivo no encontrado: {ruta_completa}")
            continue
        
        # Separación entre capítulos principales
        if idx > 1: # a partir del Capítulo II
            doc.add_page_break()
            
        procesar_archivo_md(doc, ruta_completa, es_primer_archivo=(idx == 0))
    
    # 4. Guardar archivo Word (.docx)
    print(f"\nGuardando documento Word en: {DOCX_OUTPUT}...")
    doc.save(DOCX_OUTPUT)
    print(f"[ÉXITO] Documento Word generado correctamente ({os.path.getsize(DOCX_OUTPUT)} bytes).")
    
    print("\n" + "=" * 70)
    print("PROCESO DE GENERACIÓN WORD CONCLUIDO CON ÉXITO")
    print(f"Total diagramas Mermaid procesados: {mermaid_counter}")
    print("=" * 70)

if __name__ == "__main__":
    main()
